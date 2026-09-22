import json
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import xlwt

REPO = Path(__file__).resolve().parents[3]
SCRIPTS = {
    "bm": REPO / "scripts" / "extract-bina-marga-2022.py",
    "cipta": REPO / "scripts" / "extract-cipta-karya.py",
    "hsd": REPO / "scripts" / "extract-cipta-karya-hsd.py",
    "lansekap": REPO / "scripts" / "extract-lansekap.py",
}


def run(script: Path, *args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(script), *args],
        cwd=REPO,
        capture_output=True,
        text=True,
        check=False,
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )


def test_help_names_input() -> None:
    for script in SCRIPTS.values():
        result = run(script, "--help")
        assert result.returncode == 0, result.stderr
        assert "--input" in result.stdout


def test_missing_input_exits_nonzero() -> None:
    result = run(SCRIPTS["cipta"])
    assert result.returncode != 0
    assert "/root/" not in result.stderr


def test_missing_file_exits_nonzero() -> None:
    result = run(SCRIPTS["bm"], "--input", "missing.xls")
    assert result.returncode != 0
    assert "input not found" in result.stderr
    assert "/root/" not in result.stderr


def test_cipta_fixture_is_stable(tmp_path: Path) -> None:
    source = tmp_path / "cipta.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Persiapan"
    sheet.append(["", "", "No", "Uraian", "Kode", "Satuan", "Koefisien", "Harga", "Jumlah"])
    sheet.append(["", "", "1.1.1", "Pekerjaan uji m3", None, None, None, None, 1000])
    sheet.append(["", "", "A", "Tenaga Kerja"])
    sheet.append(["", "", "1", "Pekerja", "L.01", "OH", 0.5, 100000, 50000])
    book.save(source)

    first = tmp_path / "out-a"
    second = tmp_path / "out-b"
    assert run(SCRIPTS["cipta"], "--input", str(source), "--output", str(first)).returncode == 0
    assert run(SCRIPTS["cipta"], "--input", str(source), "--output", str(second)).returncode == 0
    produced = first / "divisi-1" / "persiapan.json"
    assert produced.read_bytes() == (second / "divisi-1" / "persiapan.json").read_bytes()
    items = json.loads(produced.read_text(encoding="utf-8"))
    assert items[0]["kode_ahsp"] == "1.1.1"
    assert items[0]["tenaga_kerja"][0]["ref"] == "L.01"
    assert items[0]["tenaga_kerja"][0]["koefisien"] == 0.5


def test_hsd_fixture_is_stable(tmp_path: Path) -> None:
    source = tmp_path / "hsd.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Upah Bahan"
    sheet.append([None, None, None, "I.", "L.01", "UPAH", None, None])
    sheet.append([None, None, None, None, "L.01", "Pekerja", "OH", 100000])
    book.save(source)
    first = tmp_path / "a.json"
    second = tmp_path / "b.json"
    assert run(SCRIPTS["hsd"], "--input", str(source), "--output", str(first)).returncode == 0
    assert run(SCRIPTS["hsd"], "--input", str(source), "--output", str(second)).returncode == 0
    assert first.read_bytes() == second.read_bytes()
    payload = json.loads(first.read_text(encoding="utf-8"))
    assert payload["tenaga_kerja"][0]["ref"] == "L.01"
    assert payload["tenaga_kerja"][0]["harga_rp"] == 100000


def test_bina_marga_fixture_is_stable(tmp_path: Path) -> None:
    source = tmp_path / "bm.xls"
    book = xlwt.Workbook()
    boq = book.add_sheet("1-BOQ")
    boq.write(1, 0, "2.1.(1)")
    boq.write(1, 2, "Galian uji")
    boq.write(1, 3, "m3")
    boq.write(1, 5, 10)
    book.add_sheet("4-Basic Price")
    book.save(str(source))
    first = tmp_path / "out-a"
    second = tmp_path / "out-b"
    assert run(SCRIPTS["bm"], "--input", str(source), "--output", str(first)).returncode == 0
    assert run(SCRIPTS["bm"], "--input", str(source), "--output", str(second)).returncode == 0
    left = (first / "tenaga-kerja.json").read_bytes()
    assert left == (second / "tenaga-kerja.json").read_bytes()
    assert (first / "bahan-master.json").read_bytes() == (second / "bahan-master.json").read_bytes()


def test_lansekap_fixture_is_stable(tmp_path: Path) -> None:
    source = tmp_path / "lansekap.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Lansekap"
    sheet.append(["", "", "4.1.2.1", "Tanam pohon uji", "bh", 1, 1000, 1000, None, None, 1000])
    book.save(source)
    hsd = tmp_path / "hsd.json"
    hsd.write_text('{"bahan": []}\n', encoding="utf-8")
    first = tmp_path / "a.json"
    second_hsd = tmp_path / "hsd-b.json"
    second_hsd.write_text('{"bahan": []}\n', encoding="utf-8")
    second = tmp_path / "b.json"
    assert run(SCRIPTS["lansekap"], "--input", str(source), "--output", str(first), "--hsd-output", str(hsd)).returncode == 0
    assert run(SCRIPTS["lansekap"], "--input", str(source), "--output", str(second), "--hsd-output", str(second_hsd)).returncode == 0
    assert first.read_bytes() == second.read_bytes()
